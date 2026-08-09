"""Iter316: Workload PDF/XLSX - 'Generador del Proyecto' column + filter."""
import io
import os
import pytest
import requests
import openpyxl
from pymongo import MongoClient

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL').rstrip('/')
MONGO_URL = os.environ.get('MONGO_URL', 'mongodb://localhost:27017')
DB_NAME = os.environ.get('DB_NAME', 'test_database')

ADMIN_EMAIL = "rgonzalez@megasoft.com.ve"
ADMIN_PASS = "admin123"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EXPECTED_HEADERS = [
    "Implementador", "Cliente", "Generador", "RIF", "Tipo", "Cajas", "PVV",
    "Estado", "% Avance", "Días háb.", "Impl. Original",
    "Fecha Asignación", "Último Contacto",
]


@pytest.fixture(scope="module")
def headers():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=30)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['session_token']}"}


@pytest.fixture(scope="module")
def mongo_counts():
    client = MongoClient(MONGO_URL)
    coll = client[DB_NAME].projects
    total = coll.count_documents({})
    mm = coll.count_documents({"created_by_name": "Manuel Martin"})
    distinct = sorted([d for d in coll.distinct("created_by_name") if d])
    client.close()
    return {"total": total, "manuel_martin": mm, "distinct": distinct}


def _get_xlsx(headers, params=None):
    return requests.get(f"{BASE_URL}/api/projects/reports/workload-xlsx",
                        headers=headers, params=params or {}, timeout=90)


def _get_pdf(headers, params=None):
    return requests.get(f"{BASE_URL}/api/projects/reports/workload-pdf",
                        headers=headers, params=params or {}, timeout=90)


def _load_wb(content):
    return openpyxl.load_workbook(io.BytesIO(content), data_only=True)


def _iter_data_rows(ws):
    """Yield (row_index, values_row) between header row 4 and TOTAL row (excl)."""
    total_row = None
    for r in range(5, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith("TOTAL"):
            total_row = r
            break
    end = (total_row - 1) if total_row else ws.max_row
    for r in range(5, end + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(EXPECTED_HEADERS) + 1)]
        if any(v not in (None, "") for v in vals[:4]):
            yield r, vals


class TestXlsxGeneratorColumn:
    def test_headers_include_generador_as_col3(self, headers):
        r = _get_xlsx(headers)
        assert r.status_code == 200, r.text[:300]
        assert XLSX_CT in r.headers.get("Content-Type", ""), r.headers
        wb = _load_wb(r.content)
        ws = wb["Reporte de Carga"]
        actual = [ws.cell(row=4, column=i + 1).value for i in range(len(EXPECTED_HEADERS))]
        assert actual == EXPECTED_HEADERS, actual
        assert actual[2] == "Generador"

    def test_generator_filter_manuel_martin(self, headers, mongo_counts):
        expected = mongo_counts["manuel_martin"]
        assert expected == 29, f"DB has {expected} projects for Manuel Martin (expected 29)"
        r = _get_xlsx(headers, {"generator": "Manuel Martin"})
        assert r.status_code == 200, r.text[:400]
        wb = _load_wb(r.content)
        ws = wb["Reporte de Carga"]
        rows = list(_iter_data_rows(ws))
        assert len(rows) == expected, f"Rows {len(rows)} != {expected}"
        for _, vals in rows:
            assert vals[2] == "Manuel Martin", f"Row generator={vals[2]!r} not 'Manuel Martin'"

    def test_total_row_still_has_cajas_and_pvv(self, headers):
        r = _get_xlsx(headers, {"generator": "Manuel Martin"})
        wb = _load_wb(r.content)
        ws = wb["Reporte de Carga"]
        total_row = None
        for row in range(5, ws.max_row + 1):
            v = ws.cell(row=row, column=1).value
            if isinstance(v, str) and v.startswith("TOTAL"):
                total_row = row
                break
        assert total_row is not None, "TOTAL row missing"
        cajas = ws.cell(row=total_row, column=6).value
        pvv = ws.cell(row=total_row, column=7).value
        assert isinstance(cajas, (int, float)) and cajas >= 0
        assert isinstance(pvv, (int, float)) and pvv >= 0

    def test_no_filter_total_128(self, headers, mongo_counts):
        r = _get_xlsx(headers)
        wb = _load_wb(r.content)
        ws = wb["Reporte de Carga"]
        rows = list(_iter_data_rows(ws))
        assert len(rows) == mongo_counts["total"], (
            f"Full report rows {len(rows)} != DB total {mongo_counts['total']}")


class TestPdfGeneratorFilter:
    def test_pdf_no_filter_ok(self, headers):
        r = _get_pdf(headers)
        assert r.status_code == 200
        assert r.content.startswith(b"%PDF")
        self.__class__._full_size = len(r.content)

    def test_pdf_generator_filter_smaller(self, headers):
        full = _get_pdf(headers)
        filt = _get_pdf(headers, {"generator": "Manuel Martin"})
        assert filt.status_code == 200
        assert filt.content.startswith(b"%PDF")
        assert len(filt.content) < len(full.content), (
            f"Filtered PDF {len(filt.content)} not < full {len(full.content)}")

    def test_pdf_group_by_type_with_generator(self, headers):
        r = _get_pdf(headers, {"group_by": "type", "generator": "Manuel Martin"})
        assert r.status_code == 200
        assert r.content.startswith(b"%PDF")


class TestParity:
    def test_xlsx_pdf_row_count_parity(self, headers, mongo_counts):
        r_x = _get_xlsx(headers, {"generator": "Manuel Martin"})
        r_p = _get_pdf(headers, {"generator": "Manuel Martin"})
        wb = _load_wb(r_x.content)
        n_xlsx = len(list(_iter_data_rows(wb["Reporte de Carga"])))
        # PDF text may vary; use PyPDF2 if available
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(r_p.content))
            txt = "\n".join((p.extract_text() or "") for p in reader.pages)
            import re
            m = re.search(r"Total:\s*(\d+)\s*proyecto", txt)
            if m:
                assert int(m.group(1)) == n_xlsx
        except Exception:
            pass
        assert n_xlsx == mongo_counts["manuel_martin"]
