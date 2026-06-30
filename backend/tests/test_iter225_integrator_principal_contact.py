"""Iter225: 4 nuevos campos de Contacto Principal en Integrator.
Cubre: create (persist), validación email (422), validación interface_negotiation (422),
edit PUT (persist+survive refresh), template .xlsx (4 cols al final), bulk import (valida y rechaza)."""
import os, io, pytest, requests, csv, uuid
from openpyxl import load_workbook

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL').rstrip('/')
ADMIN = {"email": "rgonzalez@megasoft.com.ve", "password": "admin123"}
PREFIX = f"QATEST_ITER225_{uuid.uuid4().hex[:6]}_"

@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/login", json=ADMIN)
    assert r.status_code == 200, r.text
    return r.json()["session_token"]

@pytest.fixture(scope="module")
def hdr(token):
    return {"Authorization": f"Bearer {token}"}

@pytest.fixture(scope="module", autouse=True)
def cleanup(hdr):
    yield
    r = requests.get(f"{BASE_URL}/api/integrators?show_all=true", headers=hdr)
    if r.status_code == 200:
        for it in r.json():
            if (it.get("name") or "").startswith(PREFIX):
                requests.delete(f"{BASE_URL}/api/integrators/{it['integrator_id']}", headers=hdr)

def _base(name):
    return {"name": name, "integrator_type": "Integrador", "app_name": "App1",
            "integration_modality": "REST", "integrator_status": "En proceso"}

def test_create_persists_4_fields(hdr):
    payload = _base(PREFIX + "A")
    payload.update({"principal_contact_name": "Ana Pérez",
                    "principal_contact_phone": "+58 412-5551234",
                    "principal_contact_email": "ana@techpay.com",
                    "interface_negotiation": "Sí"})
    r = requests.post(f"{BASE_URL}/api/integrators", json=payload, headers=hdr)
    assert r.status_code == 200, r.text
    d = r.json()
    iid = d["integrator_id"]
    assert d["principal_contact_name"] == "Ana Pérez"
    assert d["principal_contact_phone"] == "+58 412-5551234"
    assert d["principal_contact_email"] == "ana@techpay.com"
    assert d["interface_negotiation"] == "Sí"
    # GET list shows them
    r2 = requests.get(f"{BASE_URL}/api/integrators?show_all=true", headers=hdr)
    found = [x for x in r2.json() if x["integrator_id"] == iid][0]
    assert found["principal_contact_email"] == "ana@techpay.com"
    assert found["interface_negotiation"] == "Sí"

def test_create_invalid_email_422(hdr):
    p = _base(PREFIX + "BAD_EMAIL")
    p["principal_contact_email"] = "sin_arroba.com"
    r = requests.post(f"{BASE_URL}/api/integrators", json=p, headers=hdr)
    assert r.status_code == 422, f"got {r.status_code}: {r.text}"

def test_create_invalid_negotiation_422(hdr):
    p = _base(PREFIX + "BAD_NEG")
    p["interface_negotiation"] = "Quizas"
    r = requests.post(f"{BASE_URL}/api/integrators", json=p, headers=hdr)
    assert r.status_code == 422, f"got {r.status_code}: {r.text}"

def test_edit_put_persists(hdr):
    r = requests.post(f"{BASE_URL}/api/integrators", json=_base(PREFIX + "EDIT"), headers=hdr)
    iid = r.json()["integrator_id"]
    upd = _base(PREFIX + "EDIT")
    upd.update({"principal_contact_name": "Luis Díaz",
                "principal_contact_phone": "0212-7654321",
                "principal_contact_email": "luis@x.com",
                "interface_negotiation": "No"})
    r2 = requests.put(f"{BASE_URL}/api/integrators/{iid}", json=upd, headers=hdr)
    assert r2.status_code == 200, r2.text
    # refresh: GET single
    r3 = requests.get(f"{BASE_URL}/api/integrators/{iid}", headers=hdr)
    j = r3.json()
    assert j["principal_contact_name"] == "Luis Díaz"
    assert j["principal_contact_email"] == "luis@x.com"
    assert j["interface_negotiation"] == "No"

def test_edit_put_invalid_email_422(hdr):
    r = requests.post(f"{BASE_URL}/api/integrators", json=_base(PREFIX + "EDIT2"), headers=hdr)
    iid = r.json()["integrator_id"]
    upd = _base(PREFIX + "EDIT2")
    upd["principal_contact_email"] = "no_valid"
    r2 = requests.put(f"{BASE_URL}/api/integrators/{iid}", json=upd, headers=hdr)
    assert r2.status_code == 422, r2.text

def test_template_has_4_new_cols_at_end(hdr):
    r = requests.get(f"{BASE_URL}/api/integrators/import/template", headers=hdr)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Plantilla" in wb.sheetnames
    ws = wb["Plantilla"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    expected = ["Nombre del Contacto Principal", "Teléfono Contacto Principal",
                "Email Contacto Principal", "Negociación de Interfaz"]
    assert headers[-4:] == expected, f"last 4 cols: {headers[-4:]}"

def _csv_bytes(rows, header):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8")

def test_bulk_import_valid_and_rejection(hdr):
    header = ["Nombre","Tipo","Aplicativo","Modalidad de Integración","Estatus",
              "Nombre del Contacto Principal","Teléfono Contacto Principal",
              "Email Contacto Principal","Negociación de Interfaz"]
    rows = [
        [PREFIX+"IMP_OK","Integrador","AppOK","REST","En proceso","Ana","555","ana@ok.com","Sí"],
        [PREFIX+"IMP_BAD_EMAIL","Integrador","AppB","REST","En proceso","Bad","x","correo_sin_arroba.com","No"],
        [PREFIX+"IMP_BAD_NEG","Integrador","AppC","REST","En proceso","X","y","z@z.com","Quizas"],
    ]
    files = {"file": ("imp.csv", _csv_bytes(rows, header), "text/csv")}
    r = requests.post(f"{BASE_URL}/api/integrators/import",
                      headers=hdr, files=files, data={"mode":"upsert"})
    assert r.status_code == 200, r.text
    j = r.json()
    # Good row should be created
    list_r = requests.get(f"{BASE_URL}/api/integrators?show_all=true", headers=hdr)
    names = {x["name"]: x for x in list_r.json()}
    assert PREFIX+"IMP_OK" in names
    assert names[PREFIX+"IMP_OK"]["principal_contact_email"] == "ana@ok.com"
    assert names[PREFIX+"IMP_OK"]["interface_negotiation"] == "Sí"
    # Bad rows NOT inserted
    assert PREFIX+"IMP_BAD_EMAIL" not in names, "bad email row was inserted!"
    assert PREFIX+"IMP_BAD_NEG" not in names, "bad neg row was inserted!"
    # Errors mentioning the columns
    err_cols = " ".join(e.get("column","") for e in j.get("errors", []))
    assert "Email Contacto Principal" in err_cols, j.get("errors")
    assert "Negociación de Interfaz" in err_cols or "Negociacion" in err_cols, j.get("errors")

def test_bulk_import_50_valid_rows(hdr):
    header = ["Nombre","Tipo","Aplicativo","Modalidad de Integración","Estatus",
              "Nombre del Contacto Principal","Teléfono Contacto Principal",
              "Email Contacto Principal","Negociación de Interfaz"]
    rows = []
    for i in range(50):
        rows.append([f"{PREFIX}BULK_{i:02d}","Integrador",f"App{i}","REST","En proceso",
                     f"C{i}",f"555-{i:04d}",f"c{i}@x.com","Sí" if i%2==0 else "No"])
    files = {"file": ("bulk.csv", _csv_bytes(rows, header), "text/csv")}
    r = requests.post(f"{BASE_URL}/api/integrators/import",
                      headers=hdr, files=files, data={"mode":"upsert"})
    assert r.status_code == 200, r.text
    j = r.json()
    assert (j.get("success_count",0) + j.get("updated_count",0)) >= 50, j
