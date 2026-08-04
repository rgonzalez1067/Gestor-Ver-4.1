"""Iteration 309 — Access Pay as a product in the certification matrix.
Verifies:
- GET /integrators/products includes 'Access Pay' (id 'prod_access_pay')
- GET /integrators/import/template Excel columns: AI='Access Pay', AJ='Nombre del Proyecto',
  AK='Observaciones', AL='Comercios relacionados', AM='Versión Componente',
  AN='Nombre del Contacto Principal'
- POST /integrators/import (upsert) with 'Access Pay' column value 'C' persists as
  certifications['prod_access_pay']='C' and NO 'accespay_product' field is stored.
- Integrator model no longer accepts/returns 'accespay_product'.
"""
import io
import os
import pytest
import requests
import openpyxl
from openpyxl.utils import get_column_letter

def _load_backend_url():
    v = os.environ.get("REACT_APP_BACKEND_URL")
    if not v:
        try:
            with open("/app/frontend/.env") as f:
                for line in f:
                    if line.startswith("REACT_APP_BACKEND_URL="):
                        v = line.split("=", 1)[1].strip()
                        break
        except FileNotFoundError:
            pass
    if not v:
        raise RuntimeError("REACT_APP_BACKEND_URL not set")
    return v.rstrip("/") + "/api"

BASE = _load_backend_url()
ADMIN_EMAIL = "rgonzalez@megasoft.com.ve"
ADMIN_PASS = "admin123"

TEST_PREFIX = "QAAP_"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=30)
    assert r.status_code == 200, r.text
    tok = r.json().get("session_token") or r.json().get("token")
    assert tok
    return tok


@pytest.fixture(scope="module")
def hdrs(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module", autouse=True)
def cleanup(hdrs):
    yield
    # Delete any QAAP_ test integrators created during the test run
    try:
        r = requests.get(f"{BASE}/integrators?show_all=true", headers=hdrs, timeout=30)
        if r.status_code == 200:
            for ig in r.json():
                if str(ig.get("name", "")).startswith(TEST_PREFIX):
                    iid = ig.get("integrator_id")
                    if iid:
                        requests.delete(f"{BASE}/integrators/{iid}", headers=hdrs, timeout=15)
    except Exception as e:
        print("cleanup error:", e)


def test_products_endpoint_includes_access_pay(hdrs):
    r = requests.get(f"{BASE}/integrators/products", headers=hdrs, timeout=30)
    assert r.status_code == 200, r.text
    products = r.json()
    ap = [p for p in products if p.get("service_id") == "prod_access_pay"]
    assert len(ap) == 1, f"Expected exactly one 'prod_access_pay', found: {ap}"
    assert ap[0]["name"] == "Access Pay"
    # Should be the last product
    assert products[-1]["service_id"] == "prod_access_pay"


def test_import_template_column_positions(hdrs):
    r = requests.get(f"{BASE}/integrators/import/template", headers=hdrs, timeout=30)
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Plantilla"]
    headers = {get_column_letter(c): ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)}
    print("Template headers:", headers)
    assert headers.get("AI") == "Access Pay", f"AI expected 'Access Pay' got {headers.get('AI')}"
    assert headers.get("AJ") == "Nombre del Proyecto"
    assert headers.get("AK") == "Observaciones"
    assert headers.get("AL") == "Comercios relacionados"
    assert headers.get("AM") == "Versión Componente"
    assert headers.get("AN") == "Nombre del Contacto Principal"


def _build_import_xlsx(name: str, access_pay_value: str = "C") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla"
    headers = [
        "Nombre", "Tipo", "Aplicativo", "Modalidad de Integración", "Estatus",
        "Access Pay", "Comercios relacionados", "Versión Componente",
    ]
    ws.append(headers)
    ws.append([
        name, "Integrador", "APQAAP v1", "PG Modalidad Universal", "En proceso",
        access_pay_value, "42", "v9.9.9",
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_access_pay_persists_as_certification(hdrs):
    name = f"{TEST_PREFIX}AccessPayCert"
    xlsx = _build_import_xlsx(name, "C")
    files = {"file": (f"{name}.xlsx", xlsx,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"mode": "upsert"}
    r = requests.post(f"{BASE}/integrators/import", headers=hdrs, files=files, data=data, timeout=60)
    assert r.status_code == 200, r.text
    result = r.json()
    print("import result:", {k: result.get(k) for k in ("status", "success_count", "updated_count", "errors")})
    assert (result.get("success_count", 0) + result.get("updated_count", 0)) >= 1, result

    # Fetch integrator and verify certification and absence of accespay_product
    r2 = requests.get(f"{BASE}/integrators?show_all=true", headers=hdrs, timeout=30)
    assert r2.status_code == 200
    matches = [ig for ig in r2.json() if ig.get("name") == name]
    assert matches, f"Imported integrator '{name}' not found"
    ig = matches[0]
    certs = ig.get("certifications") or {}
    assert certs.get("prod_access_pay") == "C", f"Expected certifications.prod_access_pay=='C', got {certs.get('prod_access_pay')}. Full certs: {certs}"
    assert "accespay_product" not in ig, f"'accespay_product' field must not exist on integrator: {list(ig.keys())}"
    # V3 fields still work
    assert ig.get("comercios_relacionados") == "42"
    assert ig.get("componente_version") == "v9.9.9"


def test_create_integrator_ignores_accespay_product_field(hdrs):
    """Ensure model no longer stores accespay_product even if sent."""
    payload = {
        "name": f"{TEST_PREFIX}NoAccespayField",
        "integrator_type": "Integrador",
        "app_name": "AppQAAP",
        "integration_modality": "PG Modalidad Universal",
        "integrator_status": "En proceso",
        "accespay_product": "Should be ignored",  # extra field
        "comercios_relacionados": "7",
        "componente_version": "v1.0",
    }
    r = requests.post(f"{BASE}/integrators", headers=hdrs, json=payload, timeout=30)
    # Pydantic may accept (ignore extras) or reject; both are fine as long as it's not persisted
    if r.status_code not in (200, 201):
        pytest.skip(f"Create rejected extra field (also acceptable): {r.status_code} {r.text[:200]}")
    created = r.json()
    assert "accespay_product" not in created, f"accespay_product must not be returned: {list(created.keys())}"
    assert created.get("comercios_relacionados") == "7"
    assert created.get("componente_version") == "v1.0"
