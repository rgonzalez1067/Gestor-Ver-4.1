"""Genera el MANUAL DE USUARIO DETALLADO de los 3 módulos (Cotizaciones,
Proyectos, Integradores) de la plataforma Mega Soft, en formato Word editable.

Salida: /app/frontend/public/Manual_Usuario_Detallado.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

IMG_DIR = '/tmp/manual_imgs'


def figure(fname, caption):
    """Inserta una captura de pantalla centrada con su pie de figura."""
    path = os.path.join(IMG_DIR, fname)
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(6.3))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = cap.add_run('Figura: ' + caption)
    rc.italic = True
    rc.font.size = Pt(9)
    rc.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

NAVY = RGBColor(0x14, 0x2A, 0x45)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREEN = RGBColor(0x15, 0x80, 0x3D)
AMBER = RGBColor(0xB4, 0x53, 0x09)
GREY = RGBColor(0x4B, 0x55, 0x63)

doc = Document()
base = doc.styles['Normal']
base.font.name = 'Calibri'
base.font.size = Pt(11)


def h1(t):
    p = doc.add_heading(t, level=1)
    for r in p.runs:
        r.font.color.rgb = NAVY
    return p


def h2(t):
    p = doc.add_heading(t, level=2)
    for r in p.runs:
        r.font.color.rgb = BLUE
    return p


def h3(t):
    p = doc.add_heading(t, level=3)
    for r in p.runs:
        r.font.color.rgb = NAVY
    return p


def para(t, bold=False, italic=False, color=None, size=11):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def bullet(t):
    doc.add_paragraph(t, style='List Bullet')


def step(t):
    doc.add_paragraph(t, style='List Number')


def note(kind, t):
    icons = {'tip': ('💡 Consejo: ', GREEN), 'note': ('📌 Nota: ', BLUE),
             'warn': ('⚠️ Importante: ', AMBER)}
    prefix, color = icons[kind]
    p = doc.add_paragraph()
    r = p.add_run(prefix)
    r.bold = True
    r.font.color.rgb = color
    r.font.size = Pt(10.5)
    r2 = p.add_run(t)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = color


# ============ PORTADA ============
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('Manual de Usuario'); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = NAVY
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = s.add_run('Plataforma de Gestión Comercial y Técnica — Mega Soft'); rs.font.size = Pt(14); rs.font.color.rgb = GREY
s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs2 = s2.add_run('Módulos: Cotizaciones · Proyectos · Integradores'); rs2.italic = True; rs2.font.size = Pt(12); rs2.font.color.rgb = BLUE
doc.add_paragraph()
para('Este manual describe, paso a paso, la operación diaria de los tres módulos que controlan '
     'todo el ciclo de vida comercial y técnico de nuestros aliados. Está dirigido a usuarios '
     'comerciales, técnicos/QA y gestores de aliados.', italic=True, color=GREY, size=10.5)

# Convenciones
h3('Convenciones de este manual')
bullet('Los nombres de botones y menús aparecen entre comillas, p. ej. "Crear Cotización".')
bullet('Los pasos numerados indican una secuencia que debe seguirse en orden.')
bullet('Los íconos 💡 (consejo), 📌 (nota) y ⚠️ (importante) resaltan información clave.')

# Índice
h3('Contenido')
bullet('Capítulo 1 — Menú de Cotizaciones (Usuario Comercial)')
bullet('Capítulo 2 — Menú de Proyectos (Usuario Técnico / QA)')
bullet('Capítulo 3 — Menú de Integradores (Gestión de Aliados)')
bullet('Guía de navegación rápida de la interfaz')
bullet('Glosario de términos')

doc.add_page_break()

# ==================================================================
# CAPÍTULO 1 — COTIZACIONES
# ==================================================================
h1('Capítulo 1: Menú de Cotizaciones')
para('El menú de Cotizaciones es el punto de partida de la relación comercial. Permite '
     'estructurar propuestas económicas para los clientes y darles seguimiento durante todo '
     'su ciclo de vida, desde el borrador hasta la entrega o facturación.')

h2('1.1 Acceso y vista general')
para('Al ingresar al menú verás:')
bullet('Tarjetas de indicadores (KPIs) con el conteo de cotizaciones por estado.')
bullet('Un widget de "Cotizaciones en Estado Irregular" (cuando aplica) para depurar registros.')
bullet('Botones de creación de nuevas cotizaciones.')
bullet('Una tabla central con todas las cotizaciones y sus acciones.')
bullet('Filtros de búsqueda (cliente, número, estado, tipo, segmento, fecha).')

figure('cotizaciones.jpeg', 'Menú de Cotizaciones: KPIs por estado, botones de creación '
       '(Implementaciones, Equipos y Accesorios, Reparaciones), filtros rápidos y tabla con el '
       'avance de estados por fila.')

h2('1.2 Tipos de cotización y reglas de Bienes y Servicios')
para('Al crear una cotización, el sistema filtra automáticamente los productos según el tipo '
     'de documento:')
bullet('Cotización de Equipos: venta o arrendamiento de terminales principales (ej. PinPads, POS, MPOS).')
bullet('Cotización de Accesorios: componentes complementarios (ej. cables, fuentes de poder, bases).')

h3('Regla especial de categorías combinadas')
bullet('PinPad/Accesorio (híbrido): estos productos aparecen disponibles tanto en una '
       'Cotización de Equipos como en una de Accesorios.')
bullet('PinPad/Licencia (activos lógicos / software): al cargarse en inventario el sistema '
       'NO exige seriales físicos, lo que agiliza su facturación y despacho.')
note('note', 'La categoría del producto determina en qué flujo lo verás y si requiere seriales. '
     'Si no encuentras un producto, verifica que su categoría corresponda al tipo de cotización.')

h2('1.3 Crear una cotización')
para('Desde los botones de creación puedes iniciar cuatro flujos, combinando el tipo de '
     'documento con el segmento del cliente:')
bullet('Equipos — Clientes Corporativos.')
bullet('Equipos — Clientes Pymes.')
bullet('Implementación / Proyectos de gran envergadura — Corporativos (VPOS, MPOS, Gateway, Link de Pago).')
bullet('Implementación / Proyectos — Pymes.')
para('Pasos generales:')
step('Pulsa el botón de creación correspondiente al tipo y segmento.')
step('Selecciona el cliente (buscador por Razón Social, RIF o Nombre de Fantasía).')
step('Agrega los productos/servicios con el selector; define cantidades y condiciones.')
step('Revisa los totales (Inversión Inicial, Total USD) calculados automáticamente.')
step('Guarda la cotización: quedará en estado "Borrador".')
note('tip', 'Usa el asistente (wizard) de equipos para cargar múltiples productos y sus '
     'seriales/preasignaciones de forma guiada.')

h2('1.4 La tabla de cotizaciones y sus columnas')
para('Cada fila representa una cotización. Las columnas principales son:')
bullet('Número — identificador de la cotización.')
bullet('Cliente / Nombre de Fantasía — a quién va dirigida.')
bullet('Tipo y Categoría — Equipos o Accesorios y su clasificación.')
bullet('Segmento — Corporativo o Pyme.')
bullet('Total USD / Inversión Inicial — montos de la propuesta.')
bullet('Estado — etapa actual del ciclo de vida.')
bullet('Fecha — fecha de creación o del último cambio.')
bullet('Acciones — menú con las operaciones disponibles según el estado.')
note('tip', 'Haz clic sobre una fila para abrir la ficha detallada de la cotización.')

h2('1.5 Ciclo de vida y acciones por estado')
para('Una cotización avanza por distintos estados. Las acciones disponibles en el menú '
     '"Acciones" cambian según el estado actual. Flujo típico:')
step('Borrador — la cotización recién creada; puedes editarla o enviarla al cliente.')
step('Enviada al Cliente — se remite la propuesta; queda a la espera de respuesta.')
step('Aprobada — el cliente acepta; habilita el paso a implementación/entrega.')
step('Enviada a Implementación — pasa al área técnica (genera el proyecto asociado).')
step('Entrega / Entregada — se registra la entrega física de equipos (con seriales).')
para('Otras acciones frecuentes:')
bullet('Editar / Modificar — ajusta productos o condiciones (según permisos y estado).')
bullet('Duplicar — crea una nueva a partir de una existente.')
bullet('Anular / Rechazar — cancela la propuesta dejando traza en el historial.')
bullet('Ver / Descargar PDF — genera el documento formal de la cotización.')
bullet('Preasignar seriales — reserva equipos del inventario antes de entregar.')
bullet('Regularizar — corrige cotizaciones en estado irregular (ver 1.7).')
note('note', 'Cada cambio de estado queda registrado en el "Historial" de la cotización, '
     'incluyendo excepciones y justificaciones.')

h2('1.6 Entrega, preasignación y anexos')
bullet('Preasignación: reserva seriales específicos del inventario para la cotización.')
bullet('Entrega: registra la salida física de los equipos y puede enviar un correo de entrega.')
bullet('Anexos: adjunta documentos de soporte a la cotización desde el gestor de anexos.')
note('warn', 'Los productos de categoría PinPad/Licencia no requieren seriales; el sistema '
     'omite ese paso automáticamente.')

h2('1.7 Regularización de estados irregulares')
para('Cuando una cotización queda en un estado inconsistente (por migraciones o cambios de '
     'flujo), aparece en el widget "Cotizaciones en Estado Irregular".')
step('Abre el widget de irregulares.')
step('Ejecuta primero una simulación ("dry-run") para ver qué se corregirá.')
step('Aplica la regularización individual o por lote.')
step('Revisa el resultado y confirma que el estado quedó correcto.')

h2('1.8 Migración / Respaldo de cotizaciones y anexos (ZIP)')
para('El botón de migración permite exportar/importar el paquete de cotizaciones y restaurar '
     'sus anexos desde archivos ZIP.')
bullet('Al restaurar anexos ZIP verás un monitor con el detalle y una barra de progreso '
       'porcentual con contador de restaurados/fallidos y tiempo estimado (ETA).')
note('tip', 'Puedes seleccionar uno o varios ZIP a la vez; el progreso muestra el avance global.')

h2('1.9 Regla de pertenencia y visibilidad (Seguridad de Ventas)')
para('Cada cotización se almacena con la etiqueta (badge) del departamento del creador AL '
     'MOMENTO de crearla (ej. Ventas Corporativas).')
note('warn', 'Si el usuario es promovido o transferido a otra área (ej. Implementación), sus '
     'cotizaciones históricas SIGUEN perteneciendo y siendo visibles para su equipo de Ventas '
     'original. Así el departamento de Ventas no pierde su historial ni sus métricas por '
     'movimientos de personal.')

doc.add_page_break()

# ==================================================================
# CAPÍTULO 2 — PROYECTOS
# ==================================================================
h1('Capítulo 2: Menú de Proyectos')
para('Una vez aprobada la fase comercial, el integrador inicia su proceso técnico en el Menú '
     'de Proyectos, donde se hace seguimiento al desarrollo de software e interfaces post-venta.')

h2('2.1 Panel de control y KPIs dinámicos')
para('En la parte superior verás cajas de indicadores (KPIs) que agrupan los proyectos por su '
     'estado (Por Iniciar/Asignado, En Proceso/Ejecución, Suspendidos, y el Total). También '
     'existe el filtro "Activos (ocultar cerrados)".')

h3('Visualización de avance (hover)')
para('Para conocer la salud de los proyectos sin entrar a cada uno, pasa el cursor sobre '
     'cualquier caja de KPI. Se despliega un resumen flotante ("Resumen de Avance") que detalla '
     'cuántos proyectos están:')
bullet('🟢 Al día — cumpliendo el cronograma.')
bullet('🟡 Retraso Medio — con desvíos menores en los tiempos.')
bullet('🔴 Retraso Crítico — estancados; requieren atención inmediata.')
note('tip', 'El semáforo se calcula con días hábiles, descontando fines de semana y feriados '
     'del calendario laboral configurado.')

figure('proyectos.jpeg', 'Menú de Proyectos: tableros de Avance Físico y Avance Digital · PVV, '
       'cajas KPI por estado (Total, Por Asignar, Asignado, En Gestión, Suspendido, Culminado), '
       'barra de Avance Global, filtros y tabla de seguimiento.')

h2('2.2 Columnas, filtros y búsqueda')
para('La tabla de proyectos incluye, entre otros: Cliente/Ticket, Razón Social, Sede, Tipo, '
     'Patrocinador, Implementador (Responsable Actual), Estado, Compromiso y Fecha Límite.')
bullet('Búsqueda por cliente, RIF, Razón Social o Grupo Económico.')
bullet('Filtros por "Todos los estados" y "Todos los tipos".')
bullet('Indicador de "Avance Global" y alertas de compromisos.')

h2('2.3 Asignar y reasignar proyectos')
h3('Asignar Proyecto')
step('Pulsa "Asignar Proyecto" en el proyecto correspondiente.')
step('Selecciona el implementador responsable.')
step('Confirma la asignación.')
h3('Reasignar Proyecto (individual o por lote)')
step('Usa "Reasignar Proyecto" (o "Reasignación por lote" para varios).')
step('Selecciona el nuevo implementador.')
step('Indica la Fecha de Reasignación y el Motivo (obligatorio).')
step('Confirma; el cambio queda registrado con su justificación.')

h2('2.4 Cambio de estado con justificación')
para('Al cambiar el estado de un proyecto, el sistema solicita:')
bullet('El "Nuevo Estado".')
bullet('Un "Motivo o detalle del cambio de estado" (obligatorio) para dejar trazabilidad.')

h2('2.5 Flujo de cierre de un proyecto estándar')
para('Cuando el integrador finaliza con éxito sus pruebas, ejecuta la acción de "Cierre de '
     'Proyecto". El sistema te guía por dos ventanas emergentes (modales):')
h3('Modal 1 — Datos Técnicos del Componente')
step('Transcribe el "Componente" (ej. Plugin WooCommerce, API Rest, SDK Android).')
step('Transcribe la "Versión" (ej. v2.4.1).')
note('note', 'Ambos campos son obligatorios. Con ellos el sistema arma la variable '
     '{Interfaz_Integrada} = "Componente - Versión" que se usa en el certificado y las plantillas.')
h3('Modal 2 — Personalizar Comunicación')
step('Adjunta archivos complementarios (bitácoras de QA, reportes).')
step('Ingresa correos electrónicos adicionales en copia para la notificación final.')
step('Pulsa "Cerrar y Certificar" para completar el cierre.')

h3('Excepción absoluta: Proyectos de Ambiente de Prueba')
note('warn', 'Si el proyecto es de tipo "Ambiente de Prueba", al cerrarlo el sistema ÚNICA y '
     'EXCLUSIVAMENTE lo retira de la vista de proyectos activos. No se levantan modales, no se '
     'genera certificado, no se envían correos y no afecta la grilla de integradores. Es un flujo '
     'técnico directo.')

doc.add_page_break()

# ==================================================================
# CAPÍTULO 3 — INTEGRADORES
# ==================================================================
h1('Capítulo 3: Menú de Integradores')
para('El menú de Integradores es el registro maestro y permanente de todas las empresas y '
     'comercios que han culminado o mantienen una relación técnica activa con Mega Soft.')

h2('3.1 Vista de la grilla y datos del integrador')
para('Cada fila muestra la ficha del aliado con columnas como: Nombre, Aplicativo, Categoría, '
     'Tipo de Integración, Coordinador/Gestor, Estatus (Certificado, Cerrado, En Ejecución, '
     'Suspendido, Ambiente de Prueba), Certificados y Acciones.')
bullet('Filtros por "Tipo de Integración" y por "Estatus".')
bullet('Buscador por nombre, aplicativo o gestor.')
bullet('Agrupación configurable ("Agrupar por").')

figure('integradores.jpeg', 'Menú de Integradores: barra de acciones (Resumen, Plantillas, '
       'Comunicación Masiva, Importar, Excel, PDF), filtros, cajas KPI y grilla con badges '
       '"AMPLIACIÓN" / "NUEVO COMPONENTE" y columna de Certificados.')

h2('3.2 Crear, editar, eliminar e importar')
h3('Crear / Editar integrador')
step('Pulsa "Crear Integrador" (o edita uno existente).')
step('Completa Datos del Integrador, Datos de Contacto Principal (correo, teléfono) y, si aplica, '
     'el Correo Adicional Eventual.')
step('Guarda los cambios.')
h3('Eliminar')
note('warn', 'La acción "Eliminar Integrador" es irreversible. Úsala con precaución.')
h3('Importación masiva por Excel')
step('Descarga la plantilla Excel modelo (con columnas y ejemplos).')
step('Completa los datos y súbela mediante la carga de archivo.')
step('Revisa el resultado de la importación y corrige los errores señalados.')

h2('3.3 Comportamiento post-cierre de proyectos (Inserción vs. Ampliación)')
para('La grilla se alimenta automáticamente de los cierres del menú de Proyectos:')
bullet('Nuevo Integrador o Componente: si el proyecto cerrado es de una empresa nueva o una '
       'tecnología que no tenían, se crea una fila completamente nueva.')
bullet('Ampliación: si el integrador ya existía y solo certificó una nueva etapa/actualización, '
       'el sistema reemplaza los datos de texto de su registro anterior para mantener la ficha al día.')
bullet('Coexistencia con Ambientes de Prueba: si a un integrador actual se le abre un Ambiente de '
       'Pruebas, su registro no desaparece ni se bloquea para ventas; estará visible en ambos menús '
       'simultáneamente.')

h2('3.4 Repositorio multiversión de certificados digitales')
para('En la grilla hay una columna dedicada a los Certificados Digitales.')
h3('Automatización')
para('Cada vez que cierras un proyecto estándar, el sistema toma el PDF oficial del depósito e '
     'inyecta los datos: Nombre, Aplicativo, Componente, Versión, Fecha (con la datación del cierre) '
     'y los Medios de Pago concatenados por diagonal (/). El certificado se guarda en la ficha del '
     'integrador.')
h3('Historial completo (no sobreescritura)')
step('Haz clic en el indicador de archivos (ej. "📄 Certificados (3)").')
step('Se despliega un listado cronológico con todas las versiones históricas.')
step('Desde allí puedes Descargar (📥) cada certificado.')
step('Usa el botón "+ Agregar" para subir un PDF manualmente (si tienes el permiso requerido).')
note('note', 'El sistema NUNCA borra un certificado anterior: cada etapa certificada se conserva '
     'como una versión más en el historial.')

figure('certs.jpeg', 'Repositorio multiversión de certificados: al hacer clic en el indicador '
       'se despliega el listado con opción de Descargar (📥) y el botón "+ Agregar" para subir un '
       'PDF manualmente.')

h2('3.5 Módulo de Comunicaciones Masivas (BCC)')
para('Herramienta de difusión masiva para avisar a los integradores sobre cambios en la '
     'plataforma. Ábrela con el botón "Comunicación Masiva".')
step('Segmentación: filtra por "Tipo de Integración" para ver solo a los afectados (ej. API Rest).')
step('Selección: marca la casilla maestra "Seleccionar Todos" los filtrados, o marca uno a uno.')
step('Elige una plantilla HTML preelaborada de la biblioteca de comunicación (Integradores).')
step('Adjunta archivos del repositorio (valida cada uno con 👁️ Visualizar o 📥 Descargar) '
     'o sube un archivo local desde tu computadora.')
step('Pulsa Enviar.')
note('warn', 'Privacidad obligatoria (BCC): el sistema coloca automáticamente a TODOS los '
     'integradores en Copia Oculta. Ningún aliado verá los correos ni los datos de los demás '
     'destinatarios, resguardando la confidencialidad de la base de datos.')

figure('masscomm.jpeg', 'Diálogo de Comunicación Masiva: filtro por Tipo de Integración, lista de '
       'destinatarios con casillas (maestra e individuales), selector de plantilla institucional, '
       'documentos del repositorio con Visualizar (👁️) y Descargar (📥), carga de anexo local y '
       'aviso de envío en copia oculta (BCC).')

h2('3.6 Edición de plantillas de comunicación')
bullet('Las plantillas se editan en una ventana con scroll vertical interno y pie de página fijo.')
bullet('El botón "Guardar" permanece siempre visible en la parte inferior, sin importar el largo '
       'del texto o el zoom del navegador.')

doc.add_page_break()

# ==================================================================
# NAVEGACIÓN + GLOSARIO
# ==================================================================
h1('Guía de navegación rápida de la interfaz')
bullet('Columna de "Acciones" visible: en pantallas estándar la verás a la derecha sin arrastrar '
       'la pantalla lateralmente.')
bullet('Doble barra de desplazamiento: en listas muy anchas o largas dispones de una barra de '
       'scroll horizontal arriba y abajo de la tabla, ambas sincronizadas.')
bullet('Edición segura de plantillas: scroll interno + pie fijo, con "Guardar" siempre a la vista.')
bullet('Buscadores y filtros: presentes en cada módulo para localizar registros rápidamente.')

h1('Glosario de términos')
def gl(term, desc):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(term + ': '); r.bold = True
    p.add_run(desc)
gl('KPI', 'Caja de indicador que agrupa registros por estado (proyectos o cotizaciones).')
gl('Semáforo de avance', 'Clasificación de salud del proyecto: Al día (verde), Retraso Medio (amarillo), Retraso Crítico (rojo).')
gl('Badge de departamento', 'Etiqueta inmutable que fija a qué departamento pertenece una cotización.')
gl('Componente / Versión', 'Datos técnicos capturados al cerrar un proyecto (Modal 1).')
gl('Interfaz_Integrada', 'Variable que concatena "Componente - Versión" para certificados y plantillas.')
gl('Ampliación', 'Cierre de una nueva etapa de un integrador existente; actualiza su ficha sin crear fila nueva.')
gl('Ambiente de Prueba', 'Proyecto de prueba cuyo cierre es un flujo directo (sin certificado ni correos).')
gl('BCC / Copia Oculta', 'Modo de envío donde ningún destinatario ve a los demás.')
gl('Certificado multiversión', 'Historial cronológico de certificados PDF por integrador, sin sobreescritura.')
gl('Regularización', 'Proceso para corregir cotizaciones en estado inconsistente.')

doc.add_paragraph()
para('Documento editable de referencia. Puede ajustarse la redacción, agregar capturas de '
     'pantalla y ampliar cada sección según las necesidades del equipo de adiestramiento.',
     italic=True, color=GREY, size=10)

out = '/app/frontend/public/Manual_Usuario_Detallado.docx'
doc.save(out)
print('SAVED', out)
